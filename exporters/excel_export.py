import io

from openpyxl import Workbook
from openpyxl.styles import Font


def build_excel(session: dict, proiezione: dict) -> bytes:
    wb = Workbook()
    ws_r = wb.active
    ws_r.title = "Riepilogo"
    bold = Font(bold=True)
    meta = session["meta"]
    esito = proiezione.get("esito", {})
    rows = [
        ("Comune", meta.get("comune")),
        ("Tipo elezione", proiezione.get("tipo_label", meta.get("tipo"))),
        ("Metodo proiezione", proiezione.get("metodo_label")),
        ("Avanzamento scrutinio", f"{proiezione.get('avanzamento_pct', 0)}%"),
        ("Schede scrutinate", proiezione.get("schede_scrutinate")),
        ("Schede totali", proiezione.get("schede_totali")),
        ("Voti validi (reali)", proiezione.get("voti_validi_reali")),
        ("Esito proiettato", esito.get("messaggio")),
        ("Soglia vittoria", f"{esito.get('soglia_pct', 50)}%+1"),
    ]
    for i, (k, v) in enumerate(rows, 1):
        ws_r.cell(row=i, column=1, value=k).font = bold
        ws_r.cell(row=i, column=2, value=v)

    ws_c = wb.create_sheet("Candidati")
    headers = [
        "Candidato",
        "Lista",
        "Voti reali",
        "% reali",
        "Voti proiettati",
        "% proiettate",
        "Delta soglia 50%",
    ]
    for col, h in enumerate(headers, 1):
        ws_c.cell(row=1, column=col, value=h).font = bold
    for row, c in enumerate(proiezione.get("candidati", []), 2):
        ws_c.cell(row=row, column=1, value=c.get("nome_display", c["nome"]))
        ws_c.cell(row=row, column=2, value=c.get("lista", ""))
        ws_c.cell(row=row, column=3, value=c["voti_reali"])
        ws_c.cell(row=row, column=4, value=c["pct_reali"])
        ws_c.cell(row=row, column=5, value=c["voti_proiettati"])
        ws_c.cell(row=row, column=6, value=c["pct_proiettate"])
        ws_c.cell(row=row, column=7, value=c["delta_soglia"])

    if session.get("modalita_inserimento") != "cumulativo":
        ws_s = wb.create_sheet("Sezioni")
        sh = [
            "Sezione",
            "Elettori",
            "Scrutinate",
            "Voti validi",
            "Nulli",
            "Contestati",
            "Bianche",
        ] + [c["nome"] for c in session["candidati"]]
        for col, h in enumerate(sh, 1):
            ws_s.cell(row=1, column=col, value=h).font = bold
        for row, sec in enumerate(session["sezioni"], 2):
            ws_s.cell(row=row, column=1, value=sec.get("nome"))
            ws_s.cell(row=row, column=2, value=sec["elettori"])
            ws_s.cell(row=row, column=3, value=sec.get("schede_scrutinate", 0))
            ws_s.cell(row=row, column=4, value=sec.get("voti_validi", 0))
            ws_s.cell(row=row, column=5, value=sec.get("nulli", 0))
            ws_s.cell(row=row, column=6, value=sec.get("contestati", 0))
            ws_s.cell(row=row, column=7, value=sec.get("bianche", 0))
            for ci, cand in enumerate(session["candidati"], 8):
                ws_s.cell(
                    row=row,
                    column=ci,
                    value=sec.get("voti", {}).get(cand["id"], 0),
                )

    if session.get("storico"):
        ws_h = wb.create_sheet("Storico")
        ws_h.cell(row=1, column=1, value="Timestamp").font = bold
        ws_h.cell(row=1, column=2, value="Nota").font = bold
        ws_h.cell(row=1, column=3, value="Avanzamento %").font = bold
        for row, snap in enumerate(session["storico"], 2):
            ws_h.cell(row=row, column=1, value=snap.get("timestamp"))
            ws_h.cell(row=row, column=2, value=snap.get("nota"))
            p = snap.get("proiezione", {})
            ws_h.cell(row=row, column=3, value=p.get("avanzamento_pct"))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
